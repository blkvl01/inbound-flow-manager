import unittest
from io import BytesIO
from datetime import date, datetime, timedelta
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook
import pandas as pd

import app
import data_reader


class KpiPageTests(unittest.TestCase):
    def test_time_band_hour_totals_match_original_range_masks(self):
        epoch = datetime(1899, 12, 30)
        now = datetime(2026, 6, 17, 15, 30)
        stamps = [
            datetime(2026, 6, 17, 0, 0),
            datetime(2026, 6, 17, 9, 59, 59),
            datetime(2026, 6, 17, 10, 0),
            datetime(2026, 6, 16, 23, 59),
        ]
        serials = pd.Series([(stamp - epoch).total_seconds() / 86400 for stamp in stamps] + [None])
        weights = pd.Series([1.25, 2.5, 3.75, 4.0, 99.0])
        boxes = pd.Series([1, 2, 3, 4, 99])
        parcels = pd.Series([10, 20, 30, 40, 99])
        totals = data_reader._time_band_hour_totals(serials, weights, boxes, parcels)

        for offset in (0, 1):
            bands = data_reader._time_band_inbound(
                serials, weights, now, epoch, boxes, parcels,
                day_offset=offset, hour_totals=totals,
            )
            hours = [hour for band in bands for hour in band["hours"]]
            day = now.replace(hour=0, minute=0, second=0, microsecond=0) - timedelta(days=offset)
            for hour in hours:
                start = day + timedelta(hours=hour["start"])
                end = start + timedelta(hours=1)
                mask = serials.between(
                    (start - epoch).total_seconds() / 86400,
                    (end - epoch).total_seconds() / 86400,
                    inclusive="left",
                )
                self.assertEqual(hour["count"], int(mask.sum()))
                self.assertEqual(hour["kg"], round(float(weights[mask].sum()), 1))
                self.assertEqual(hour["colli"], round(float(boxes[mask].sum()), 1))
                self.assertEqual(hour["parcel"], round(float(parcels[mask].sum()), 1))

    def test_kpi_page_payload_contains_inbound_and_outbound_shift_series(self):
        kpi = {
            "shift_name": "Jelenlegi",
            "shift_range": "06:00-14:00",
            "shift_hourly": [{"label": "06", "kg": 100}, {"label": "07", "kg": 200}],
            "shift_prev_name": "Előző",
            "shift_prev_range": "22:00-06:00",
            "shift_prev_hourly": [{"label": "22", "kg": 50}],
            "outbound_shift_name": "Outbound most",
            "outbound_shift_range": "06:00-14:00",
            "outbound_shift_hourly": [{"label": "06", "kg": 80}],
            "outbound_shift_prev_name": "Outbound előző",
            "outbound_shift_prev_range": "22:00-06:00",
            "outbound_shift_prev_hourly": [{"label": "22", "kg": 20}],
        }

        payload = app._kpi_page_payload(kpi, datetime(2026, 6, 15, 10, 5))

        self.assertEqual(payload["updated"], "06.15 10:05")
        self.assertEqual(payload["inbound"]["current"]["kg"], 300)
        self.assertEqual(payload["inbound"]["prev"]["kg"], 50)
        self.assertEqual(payload["outbound"]["current"]["kg"], 80)
        self.assertEqual(payload["outbound"]["prev"]["range"], "22:00-06:00")

    def test_kpi_page_button_keeps_original_hungarian_ops_label(self):
        text, class_name = app.update_kpi_page_button("kpi")

        self.assertEqual(text, "Operatív")
        self.assertIn("kpi-page-btn-active", class_name)

    def test_kpi_page_payload_preserves_tracking_section(self):
        kpi = {
            "tracking": {
                "current_warehouse_saturation": {"kg": 125000, "ratio": 0.25, "capacity": 500000},
                "expected_arrivals": [{"date": "2026-06-15", "expected_arrivals_kg": 100}],
            }
        }

        payload = app._kpi_page_payload(kpi, datetime(2026, 6, 15, 10, 5))

        self.assertEqual(payload["tracking"]["current_warehouse_saturation"]["kg"], 125000)
        self.assertEqual(payload["tracking"]["expected_arrivals"][0]["date"], "2026-06-15")

    def test_time_band_payload_includes_colli_and_parcel_metrics(self):
        payload = data_reader._merge_time_bands(
            [{
                "key": "08-12",
                "label": "08-12",
                "start": 8,
                "end": 12,
                "kg": 120.5,
                "count": 2,
                "colli": 7,
                "parcel": 34,
                "active": True,
                "current": True,
            }],
            [{
                "key": "08-12",
                "kg": 80,
                "count": 1,
                "colli": 5,
                "parcel": 13,
            }],
            datetime(2026, 6, 17, 10, 0),
        )

        bands = {band["key"]: band for band in payload["bands"]}
        band = bands["08-12"]
        self.assertEqual(payload["date"], "2026-06-17")
        self.assertEqual(band["inbound_colli"], 7)
        self.assertEqual(band["inbound_parcel"], 34)
        self.assertEqual(band["outbound_colli"], 5)
        self.assertEqual(band["outbound_parcel"], 13)
        self.assertEqual(len(band["hours"]), 4)

    def test_time_band_outbound_uses_ecomm_departure_serials(self):
        epoch = datetime(1899, 12, 30)
        now = datetime(2026, 6, 17, 15, 30)

        def serial(dt):
            return (dt - epoch).total_seconds() / 86400

        departure = pd.Series([
            serial(datetime(2026, 6, 17, 9, 15)),
            serial(datetime(2026, 6, 17, 13, 5)),
            serial(datetime(2026, 6, 16, 13, 5)),
        ])
        weights = pd.Series([100.0, 200.0, 999.0])
        boxes = pd.Series([2.0, 3.0, 99.0])
        parcels = pd.Series([10.0, 20.0, 999.0])

        outbound = data_reader._time_band_outbound_ecomm(departure, weights, now, epoch, boxes, parcels)
        inbound = data_reader._time_band_inbound(
            pd.Series(dtype=float), pd.Series(dtype=float), now, epoch
        )
        payload = data_reader._merge_time_bands(inbound, outbound, now)
        bands = {band["key"]: band for band in payload["bands"]}
        hours = {hour["key"]: hour for band in payload["bands"] for hour in band["hours"]}

        self.assertEqual(bands["08-12"]["outbound_kg"], 100.0)
        self.assertEqual(bands["08-12"]["outbound_colli"], 2.0)
        self.assertEqual(hours["09-10"]["outbound_parcel"], 10.0)
        self.assertEqual(bands["12-16"]["outbound_kg"], 200.0)
        self.assertEqual(bands["12-16"]["outbound_parcel"], 20.0)

    def test_time_band_payload_can_include_all_available_history_days(self):
        epoch = datetime(1899, 12, 30)
        now = datetime(2026, 6, 17, 15, 30)

        def serial(dt):
            return (dt - epoch).total_seconds() / 86400

        serials = pd.Series([
            serial(datetime(2026, 6, 17, 9, 15)),
            serial(datetime(2026, 6, 15, 13, 5)),
            serial(datetime(2026, 6, 12, 10, 0)),
        ])
        weights = pd.Series([100.0, 200.0, 300.0])
        boxes = pd.Series([1.0, 2.0, 3.0])
        parcels = pd.Series([10.0, 20.0, 30.0])

        offsets = data_reader._time_band_day_offsets(now, epoch, serials)
        inbound_days = {
            offset: data_reader._time_band_inbound(serials, weights, now, epoch, boxes, parcels, day_offset=offset)
            for offset in offsets
        }
        payload = data_reader._merge_time_bands(inbound_days, {}, now)

        self.assertEqual([day["date"] for day in payload["days"]], [
            "2026-06-12",
            "2026-06-15",
            "2026-06-17",
        ])
        self.assertEqual(payload["today_index"], 2)
        by_day = {day["date"]: day for day in payload["days"]}
        june_12_bands = {band["key"]: band for band in by_day["2026-06-12"]["bands"]}
        june_15_bands = {band["key"]: band for band in by_day["2026-06-15"]["bands"]}
        self.assertEqual(june_12_bands["08-12"]["inbound_kg"], 300.0)
        self.assertEqual(june_15_bands["12-16"]["inbound_parcel"], 20.0)

    def test_temu_stage_payload_includes_country_and_lmp_scopes(self):
        epoch = datetime(1899, 12, 30)
        now = datetime(2026, 6, 17, 15, 30)

        def serial(dt):
            return (dt - epoch).total_seconds() / 86400

        raw = pd.DataFrame([
            {
                "lmp": "TEMU CZ PACK",
                "ak_raw": serial(datetime(2026, 6, 17, 8, 0)),
                "noa_raw": serial(datetime(2026, 6, 17, 10, 0)),
                "am_raw": serial(datetime(2026, 6, 17, 13, 30)),
                "aq_raw": serial(datetime(2026, 6, 17, 15, 0)),
                "ar_raw": serial(datetime(2026, 6, 17, 16, 0)),
                "departure_raw": serial(datetime(2026, 6, 18, 8, 0)),
            },
            {
                "lmp": "TEMU HU PACK",
                "ak_raw": serial(datetime(2026, 6, 16, 9, 0)),
                "noa_raw": serial(datetime(2026, 6, 16, 10, 0)),
                "am_raw": serial(datetime(2026, 6, 16, 11, 0)),
                "aq_raw": serial(datetime(2026, 6, 16, 12, 0)),
                "ar_raw": serial(datetime(2026, 6, 16, 13, 0)),
                "departure_raw": serial(datetime(2026, 6, 16, 15, 0)),
            },
            {
                "lmp": "TEMU MD PACK",
                "ak_raw": serial(datetime(2026, 6, 17, 8, 0)),
                "noa_raw": serial(datetime(2026, 6, 17, 9, 0)),
                "am_raw": serial(datetime(2026, 6, 17, 10, 0)),
                "aq_raw": serial(datetime(2026, 6, 17, 11, 0)),
                "ar_raw": serial(datetime(2026, 6, 17, 12, 0)),
                "departure_raw": serial(datetime(2026, 6, 17, 13, 0)),
            },
            {
                "lmp": "NOT TEMU",
                "ak_raw": serial(datetime(2026, 6, 17, 8, 0)),
                "noa_raw": serial(datetime(2026, 6, 17, 9, 0)),
                "am_raw": serial(datetime(2026, 6, 17, 10, 0)),
                "aq_raw": serial(datetime(2026, 6, 17, 11, 0)),
                "ar_raw": serial(datetime(2026, 6, 17, 12, 0)),
                "departure_raw": serial(datetime(2026, 6, 17, 13, 0)),
            },
        ])

        payload = data_reader._compute_temu_stage_kpi(raw, now, epoch)

        self.assertEqual(payload["countries"], ["CZ", "HU"])
        self.assertEqual(payload["lmps"], ["TEMU CZ PACK", "TEMU HU PACK"])
        self.assertNotIn("TEMU MD PACK", payload["lmps"])

        cz_day = payload["day"]["CZ"][0]
        self.assertEqual(cz_day["period_key"], "2026-06-17")
        self.assertTrue(cz_day["is_today"])
        self.assertEqual(cz_day["count"], 1)
        self.assertEqual(cz_day["segments"]["ata_noa"], 2.0)
        self.assertEqual(cz_day["segments"]["noa_transfer"], 3.5)
        self.assertEqual(cz_day["segments"]["outbound"], 16.0)
        self.assertEqual(cz_day["seg_sum"]["outbound"], 16.0)
        self.assertEqual(cz_day["seg_cnt"]["outbound"], 1)

        lmp_day = payload["day"]["TEMU CZ PACK"][0]
        self.assertEqual(lmp_day["period_key"], cz_day["period_key"])
        self.assertEqual(lmp_day["total"], cz_day["total"])

    def test_temu_stage_week_payload_marks_full_week_coverage(self):
        epoch = datetime(1899, 12, 30)
        now = datetime.combine(date.fromisocalendar(2026, 25, 1), datetime.min.time())

        def serial(dt):
            return (dt - epoch).total_seconds() / 86400

        rows = []
        for iso_week, day_count in ((23, 7), (24, 2)):
            monday = date.fromisocalendar(2026, iso_week, 1)
            for offset in range(day_count):
                base = datetime.combine(monday + timedelta(days=offset), datetime.min.time()).replace(hour=8)
                rows.append({
                    "lmp": "TEMU HU PACK",
                    "ak_raw": serial(base),
                    "noa_raw": serial(base + timedelta(hours=1)),
                    "am_raw": serial(base + timedelta(hours=2)),
                    "aq_raw": serial(base + timedelta(hours=3)),
                    "ar_raw": serial(base + timedelta(hours=4)),
                    "departure_raw": serial(base + timedelta(hours=5)),
                })

        payload = data_reader._compute_temu_stage_kpi(pd.DataFrame(rows), now, epoch)
        weeks = {w["period_key"]: w for w in payload["week"]["ALL"]}

        self.assertTrue(weeks["2026-W23"]["complete"])
        self.assertEqual(weeks["2026-W23"]["day_count"], 7)
        self.assertEqual(weeks["2026-W23"]["expected_days"], 7)
        self.assertFalse(weeks["2026-W24"]["complete"])
        self.assertEqual(weeks["2026-W24"]["day_count"], 2)

    def test_temu_stage_payload_buckets_on_noa_day(self):
        epoch = datetime(1899, 12, 30)
        now = datetime(2026, 6, 18, 15, 30)

        def serial(dt):
            return (dt - epoch).total_seconds() / 86400

        raw = pd.DataFrame([{
            "lmp": "TEMU HU PACK",
            "ak_raw": serial(datetime(2026, 6, 14, 22, 0)),
            "noa_raw": serial(datetime(2026, 6, 15, 1, 0)),
            "am_raw": serial(datetime(2026, 6, 15, 3, 0)),
            "aq_raw": serial(datetime(2026, 6, 15, 4, 0)),
            "ar_raw": serial(datetime(2026, 6, 15, 5, 0)),
            "departure_raw": serial(datetime(2026, 6, 15, 8, 0)),
        }])

        payload = data_reader._compute_temu_stage_kpi(raw, now, epoch)

        self.assertEqual(payload["day"]["ALL"][0]["period_key"], "2026-06-15")
        self.assertEqual(payload["week"]["ALL"][0]["period_key"], "2026-W25")

    def test_temu_stage_guard_fixes_offbyone_typo(self):
        # A single off-by-one ATA slip (entered as tomorrow) is a trusted, unambiguous
        # fix: the date is corrected back and the shipment buckets on the right day.
        epoch = datetime(1899, 12, 30)
        now = datetime(2026, 6, 17, 15, 30)

        def serial(dt):
            return (dt - epoch).total_seconds() / 86400

        raw = pd.DataFrame([{
            "lmp": "TEMU HU PACK",
            "ak_raw": serial(datetime(2026, 6, 14, 8, 0)),   # typo: +1 day
            "noa_raw": serial(datetime(2026, 6, 13, 9, 0)),
            "am_raw": serial(datetime(2026, 6, 13, 10, 0)),
            "aq_raw": serial(datetime(2026, 6, 13, 11, 0)),
            "ar_raw": serial(datetime(2026, 6, 13, 12, 0)),
            "departure_raw": serial(datetime(2026, 6, 13, 14, 0)),
        }])

        payload = data_reader._compute_temu_stage_kpi(raw, now, epoch)
        day = payload["day"]["HU"][0]

        self.assertEqual(day["period_key"], "2026-06-13")          # NOA/report day
        self.assertEqual(day["segments"]["ata_noa"], 1.0)
        self.assertEqual(day["segments"]["customs"], 1.0)
        self.assertEqual(day["segments"]["outbound"], 2.0)
        self.assertGreaterEqual(payload["quality"]["corrected_fields"], 1)
        self.assertEqual(payload["quality"]["blocked_gaps"], 0)
        self.assertEqual(payload["suspects"]["total"], 0)

    def test_temu_stage_guard_flags_unreliable_milestone_instead_of_fabricating(self):
        # A clearly-garbage customs-end (7 days off, after departure, in the future)
        # must NOT be "fixed" by relocating the real departure. The stage is blocked,
        # the shipment stays in the KPI through its good stages, and the bad field is
        # surfaced as a flagged record for the operator to correct at source.
        epoch = datetime(1899, 12, 30)
        now = datetime(2026, 6, 17, 15, 30)

        def serial(dt):
            return (dt - epoch).total_seconds() / 86400

        raw = pd.DataFrame([{
            "awb": "999-12345678",
            "lmp": "TEMU AT AP",
            "ak_raw": serial(datetime(2026, 6, 13, 8, 0)),
            "noa_raw": serial(datetime(2026, 6, 13, 9, 0)),
            "am_raw": serial(datetime(2026, 6, 13, 10, 0)),
            "aq_raw": serial(datetime(2026, 6, 13, 11, 0)),
            "ar_raw": serial(datetime(2026, 6, 20, 12, 0)),   # garbage: future + after departure
            "departure_raw": serial(datetime(2026, 6, 13, 14, 0)),
        }])

        payload = data_reader._compute_temu_stage_kpi(raw, now, epoch)
        day = payload["day"]["AT"][0]

        # Earlier, reliable stages still count; the shipment is in the KPI.
        self.assertEqual(day["count"], 1)
        self.assertEqual(day["segments"]["ata_noa"], 1.0)
        self.assertEqual(day["segments"]["transfer_customs"], 1.0)
        # The two stages touching the garbage customs-end are blocked, never fabricated.
        self.assertEqual(day["segments"]["customs"], 0.0)
        self.assertEqual(day["segments"]["outbound"], 0.0)
        self.assertEqual(day["seg_cnt"]["customs"], 0)
        self.assertEqual(day["seg_cnt"]["outbound"], 0)
        self.assertGreaterEqual(payload["quality"]["blocked_gaps"], 2)
        self.assertEqual(payload["quality"]["corrected_fields"], 0)   # nothing relocated

        # The unreliable record is collected for the operator.
        self.assertEqual(payload["suspects"]["total"], 1)
        item = payload["suspects"]["items"][0]
        self.assertEqual(item["awb"], "999-12345678")
        self.assertEqual([f["key"] for f in item["fields"]], ["ar_raw"])
        self.assertEqual(item["fields"][0]["label"], "Customs end")
        self.assertIn("06-20", item["fields"][0]["raw"])
        self.assertEqual([m["key"] for m in item["milestones"]], [
            "ak_raw", "noa_raw", "am_raw", "aq_raw", "ar_raw", "departure_raw",
        ])
        self.assertEqual(item["milestones"][3]["iso"], "2026-06-13T11:00")
        self.assertEqual(item["milestones"][4]["iso"], "2026-06-20T12:00")
        self.assertEqual(item["milestones"][5]["iso"], "2026-06-13T14:00")

    def test_temu_stage_date_guard_repairs_swapped_day_month(self):
        epoch = datetime(1899, 12, 30)
        now = datetime(2026, 6, 17, 15, 30)

        def serial(dt):
            return (dt - epoch).total_seconds() / 86400

        raw = pd.DataFrame([{
            "lmp": "TEMU CZ PACK",
            "ak_raw": serial(datetime(2026, 6, 12, 8, 0)),
            "noa_raw": serial(datetime(2026, 12, 6, 10, 0)),
            "am_raw": serial(datetime(2026, 6, 12, 11, 0)),
            "aq_raw": serial(datetime(2026, 6, 12, 12, 0)),
            "ar_raw": serial(datetime(2026, 6, 12, 13, 0)),
            "departure_raw": serial(datetime(2026, 6, 12, 16, 0)),
        }])

        payload = data_reader._compute_temu_stage_kpi(raw, now, epoch)
        day = payload["day"]["CZ"][0]

        self.assertEqual(day["period_key"], "2026-06-12")
        self.assertEqual(day["segments"]["ata_noa"], 2.0)
        self.assertEqual(day["segments"]["noa_transfer"], 1.0)
        self.assertGreaterEqual(payload["quality"]["corrected_fields"], 1)
        self.assertEqual(payload["quality"]["blocked_gaps"], 0)

    def test_temu_stage_date_guard_blocks_uncorrectable_future_gaps(self):
        epoch = datetime(1899, 12, 30)
        now = datetime(2026, 6, 17, 15, 30)

        def serial(dt):
            return (dt - epoch).total_seconds() / 86400

        raw = pd.DataFrame([{
            "lmp": "TEMU HU PACK",
            "ak_raw": serial(datetime(2026, 6, 1, 8, 0)),
            "noa_raw": serial(datetime(2026, 6, 1, 9, 0)),
            "am_raw": serial(datetime(2026, 6, 1, 10, 0)),
            "aq_raw": serial(datetime(2026, 6, 1, 11, 0)),
            "ar_raw": serial(datetime(2026, 9, 1, 12, 0)),
            "departure_raw": serial(datetime(2026, 9, 1, 13, 0)),
        }])

        payload = data_reader._compute_temu_stage_kpi(raw, now, epoch)
        day = payload["day"]["HU"][0]

        self.assertEqual(day["segments"]["customs"], 0.0)
        self.assertEqual(day["segments"]["outbound"], 0.0)
        self.assertEqual(day["seg_cnt"]["customs"], 0)
        self.assertEqual(day["seg_cnt"]["outbound"], 0)
        self.assertGreaterEqual(payload["quality"]["blocked_gaps"], 2)

    def test_kpi_trend_export_builds_xlsx_from_visible_payload(self):
        payload = {
            "title": "Inbound trend",
            "mode": "day",
            "modeLabel": "Daily",
            "breakdown": "country",
            "breakdownLabel": "Countries",
            "rangeLabel": "W23",
            "generatedAt": "2026-06-16T12:00:00",
            "columns": [
                {"id": "HU", "label": "HU · Hungary", "color": "#34d399"},
                {"id": "AT", "label": "AT · Austria", "color": "#ff8a3d"},
            ],
            "rows": [
                {"key": "06.01", "d0": "2026-06-01", "d1": "2026-06-01", "vals": {"HU": 1200, "AT": 3400}},
                {"key": "06.02", "d0": "2026-06-02", "d1": "2026-06-02", "vals": {"HU": 2200, "AT": 1400}},
            ],
        }

        content, filename = app._build_kpi_trend_export(payload)
        wb = load_workbook(BytesIO(content), data_only=False)
        ws = wb["Data"]

        self.assertTrue(filename.endswith(".xlsx"))
        self.assertEqual(ws["A1"].value, "W23 · Daily · Countries")
        self.assertEqual(ws["A2"].value, "Export: 2026-06-16 12:00")
        self.assertNotIn("T12:00:00", ws["A2"].value)
        self.assertEqual(
            [ws.cell(3, col).value for col in range(1, 6)],
            ["Day", "Date", "HU · Hungary", "AT · Austria", "Total"],
        )
        self.assertEqual(ws.cell(4, 3).value, 1200)
        self.assertEqual(ws.cell(5, 5).value, 3600)
        self.assertTrue(ws._charts)
        self.assertIn("Summary", wb.sheetnames)
        with ZipFile(BytesIO(content)) as zf:
            self.assertTrue(any(name.startswith("xl/charts/") for name in zf.namelist()))

    def test_kpi_temu_export_uses_excel_dates_and_duration_formats(self):
        payload = {
            "kind": "table",
            "title": "TEMU KPI",
            "scopeLabel": "All countries",
            "modeLabel": "Daily",
            "periodType": "day",
            "rangeLabel": "06.13",
            "stages": [
                {"key": "ata_noa", "label": "ATA-NOA", "color": "FF0000"},
                {"key": "customs", "label": "Customs Clearance", "color": "00B0F0"},
            ],
            "rows": [{
                "label": "06.13",
                "date": "2026-06-13",
                "dateEnd": "2026-06-13",
                "count": 2,
                "vals": {"ata_noa": 1.5, "customs": 25.25},
                "total": 26.75,
            }],
            "totals": {"vals": {"ata_noa": 1.5, "customs": 25.25}, "total": 26.75, "count": 2},
        }

        content, filename = app._build_kpi_temu_export(payload)
        wb = load_workbook(BytesIO(content), data_only=False)
        ws = wb["TEMU KPI"]

        self.assertTrue(filename.endswith(".xlsx"))
        self.assertEqual(
            [ws.cell(4, col).value for col in range(1, 6)],
            ["Date", "ATA-NOA", "Customs Clearance", "Total lead time", "Pcs"],
        )
        self.assertEqual(ws.cell(5, 1).value.date(), date(2026, 6, 13))
        self.assertEqual(ws.cell(5, 1).number_format, "yyyy-mm-dd")
        self.assertEqual(ws.cell(5, 2).value, timedelta(minutes=90))
        self.assertEqual(ws.cell(5, 2).number_format, "[h]:mm")
        self.assertEqual(ws.cell(5, 4).value, timedelta(hours=26, minutes=45))
        self.assertEqual(ws.cell(5, 4).number_format, "[h]:mm")

    def test_kpi_temu_flagged_export_writes_full_milestone_chain(self):
        # The flagged-records export mirrors the on-screen "Flagged records" drawer
        # ("full milestone chain"): one row per shipment with identity columns
        # (Country, AWB, LMP, Anchor date) followed by the whole ATA→…→Departure
        # chain, the flagged step highlighted with a comment. (This superseded the
        # earlier neighbour-context-only layout.)
        payload = {
            "kind": "flagged",
            "title": "TEMU flagged records",
            "scopeLabel": "AT · Austria",
            "modeLabel": "Daily",
            "periodType": "day",
            "rangeLabel": "06.13",
            "items": [{
                "awb": "999-12345678",
                "lmp": "TEMU AT AP",
                "country": "AT",
                "day": "2026-06-13",
                "fields": [{
                    "key": "aq_raw",
                    "label": "Customs start",
                    "reason": "before previous step",
                    "iso": "2026-06-12T11:00",
                }],
                "milestones": [
                    {"key": "ak_raw", "label": "ATA", "iso": "2026-06-13T08:00"},
                    {"key": "noa_raw", "label": "NOA", "iso": "2026-06-13T09:00"},
                    {"key": "am_raw", "label": "Transfer", "iso": "2026-06-13T10:00"},
                    {"key": "aq_raw", "label": "Customs start", "iso": "2026-06-12T11:00",
                     "suspect": True, "reason": "before previous step"},
                    {"key": "ar_raw", "label": "Customs end", "iso": "2026-06-13T12:00"},
                    {"key": "departure_raw", "label": "Departure", "iso": "2026-06-13T14:00"},
                ],
            }],
        }

        content, filename = app._build_kpi_temu_export(payload)
        wb = load_workbook(BytesIO(content), data_only=False)
        ws = wb["Flagged records"]

        self.assertTrue(filename.endswith(".xlsx"))
        # Header row (4): identity columns + the full milestone chain.
        self.assertEqual(
            [ws.cell(4, c).value for c in range(1, 11)],
            ["Country", "AWB", "LMP", "Anchor date",
             "ATA", "NOA", "Transfer", "Customs start", "Customs end", "Departure"],
        )
        # Data row (5): identity then every milestone date in order.
        self.assertEqual(ws.cell(5, 1).value, "AT")
        self.assertEqual(ws.cell(5, 2).value, "999-12345678")
        self.assertEqual(ws.cell(5, 3).value, "TEMU AT AP")
        self.assertEqual(ws.cell(5, 4).value.date(), date(2026, 6, 13))
        self.assertEqual(ws.cell(5, 4).number_format, "yyyy-mm-dd")
        self.assertEqual(ws.cell(5, 5).value, datetime(2026, 6, 13, 8, 0))    # ATA
        self.assertEqual(ws.cell(5, 7).value, datetime(2026, 6, 13, 10, 0))   # Transfer
        self.assertEqual(ws.cell(5, 8).value, datetime(2026, 6, 12, 11, 0))   # Customs start (flagged)
        self.assertEqual(ws.cell(5, 9).value, datetime(2026, 6, 13, 12, 0))   # Customs end
        self.assertEqual(ws.cell(5, 10).value, datetime(2026, 6, 13, 14, 0))  # Departure
        self.assertEqual(ws.cell(5, 5).number_format, "yyyy-mm-dd hh:mm")
        self.assertEqual(ws.cell(5, 8).number_format, "yyyy-mm-dd hh:mm")
        # The out-of-order milestone is flagged in-place with a date-guard comment.
        self.assertIsNotNone(ws.cell(5, 8).comment)
        self.assertIn("Flagged", ws.cell(5, 8).comment.text)

    def test_outbound_ecomm_glabs_are_excluded_everywhere_outbound(self):
        now = datetime.now()
        pallets = pd.DataFrame([
            {
                "bud_lmp": "TEMU",
                "awb": "E1",
                "plate": "AAA-111",
                "glabs_id": "GLABS-E",
                "boxes": 1,
                "weight": 100,
                "driver_checkin": now,
                "issued_time": now,
                "rest_until": None,
                "rest_declared": False,
                "is_ecomm": True,
                "kiad_status": "Kiadva",
                "pallets_issued": 1,
                "location": "",
                "load_note": "",
            },
            {
                "bud_lmp": "TEMU",
                "awb": "E2",
                "plate": "AAA-111",
                "glabs_id": "GLABS-E",
                "boxes": 1,
                "weight": 200,
                "driver_checkin": now,
                "issued_time": now,
                "rest_until": None,
                "rest_declared": False,
                "is_ecomm": False,
                "kiad_status": "Kiadva",
                "pallets_issued": 1,
                "location": "",
                "load_note": "",
            },
            {
                "bud_lmp": "TEMU",
                "awb": "N1",
                "plate": "BBB-222",
                "glabs_id": "GLABS-N",
                "boxes": 1,
                "weight": 300,
                "driver_checkin": now,
                "issued_time": now,
                "rest_until": None,
                "rest_declared": False,
                "is_ecomm": False,
                "kiad_status": "Kiadva",
                "pallets_issued": 1,
                "location": "",
                "load_note": "",
            },
        ])

        kpi = data_reader._compute_outbound_shift_kpi(pallets)
        cards, errors = data_reader.load_outbound_data(pallets=pallets)

        self.assertEqual(kpi["outbound_shift_kg"], 300)
        self.assertEqual(kpi["outbound_shift_count"], 1)
        self.assertFalse(errors)
        self.assertEqual([card["plate"] for card in cards], ["BBB-222"])

    def test_outbound_shift_big_average_uses_completed_shift_windows(self):
        now = datetime.now()
        current_start, _, _ = data_reader._current_shift_window(now)
        shift_a = current_start - timedelta(hours=12)
        shift_b = current_start - timedelta(hours=24)

        rows = []

        def add_row(shift_start, offset_h, weight, glabs, plate):
            rows.append({
                "bud_lmp": "TEMU",
                "awb": f"{glabs}-{offset_h}",
                "plate": plate,
                "glabs_id": glabs,
                "boxes": 1,
                "weight": weight,
                "driver_checkin": shift_start + timedelta(minutes=15),
                "issued_time": shift_start + timedelta(hours=offset_h),
                "rest_until": None,
                "rest_declared": False,
                "is_ecomm": False,
                "kiad_status": "Kiadva",
                "pallets_issued": 1,
                "location": "",
                "load_note": "",
            })

        add_row(shift_a, 1, 6000, "A-1", "AAA-111")
        add_row(shift_a, 2, 6000, "A-2", "AAA-111")
        add_row(shift_b, 1, 6000, "B-1", "BBB-222")
        add_row(shift_b, 2, 6000, "B-2", "BBB-222")
        add_row(shift_b, 3, 6000, "B-3", "BBB-222")
        add_row(shift_b, 4, 6000, "B-4", "BBB-222")
        add_row(current_start, 1, 999999, "NOW-1", "NOW-999")

        kpi = data_reader._compute_outbound_shift_kpi(pd.DataFrame(rows))

        self.assertEqual(kpi["outbound_shift_avg_shift_count"], 2)
        self.assertAlmostEqual(kpi["outbound_shift_avg_kg_per_h"], 1500.0)
        self.assertAlmostEqual(kpi["outbound_shift_avg_loadings_per_h"], 0.25)

    def test_tv_ops_report_issued_count_uses_current_shift_kpi(self):
        closed_cards = [
            {"plate": f"CLOSED-{i:02d}", "active": 0, "weight": 1000, "glabs_items": []}
            for i in range(19)
        ]
        state = {
            "status": "ready",
            "last_refresh": datetime(2026, 7, 2, 9, 30),
            "df": pd.DataFrame(),
            "outbound_cards": closed_cards,
            "kpi": {
                "shift_kg": 1200,
                "shift_count": 2,
                "shift_name": "Nappali műszak",
                "shift_range": "08:00–20:00",
                "shift_elapsed_h": 2.0,
                "outbound_shift_kg": 3000,
                "outbound_shift_count": 7,
                "outbound_shift_colli": 50,
                "outbound_shift_pallets": 6,
                "outbound_shift_parcels": 80,
                "outbound_shift_trucks": 3,
                "outbound_shift_loadings": 4,
                "outbound_shift_loading_time_min": 30,
                "outbound_shift_loading_time_count": 4,
                "outbound_shift_name": "Nappali műszak",
                "outbound_shift_range": "08:00–20:00",
                "outbound_shift_hourly": [],
                "outbound_shift_prev_same_kg": 0,
                "outbound_shift_prev_same_loadings": 0,
                "outbound_shift_prev_same_hourly": [],
                "outbound_shift_avg_kg_per_h": 1500,
                "outbound_shift_avg_loadings_per_h": 2.5,
                "outbound_shift_avg_shift_count": 12,
            },
        }

        with patch.object(app.data_cache, "get_state", return_value=state):
            payload = app.server.test_client().get("/api/tv/ops").get_json()

        self.assertEqual(payload["issued_trucks"], 3)
        self.assertEqual(payload["shift_out_kg"], 3000)
        self.assertEqual(payload["report"]["out"]["issued_trucks"], 3)
        self.assertEqual(payload["report"]["out"]["trucks"], 3)
        self.assertEqual(payload["report"]["out"]["ref_kg_kind"], "big_avg")
        self.assertEqual(payload["report"]["out"]["ref_kg_per_h"], 1500)
        self.assertEqual(payload["report"]["out"]["ref_loadings_kind"], "big_avg")
        self.assertEqual(payload["report"]["out"]["ref_loadings_per_h"], 2.5)

    def test_tv_loading_note_parser_recognizes_only_safe_ramp_and_name_patterns(self):
        cases = {
            "16 RAMP / Alfonz": ("16", ["Alfonz"]),
            "16-0S RÁMPÁRA": ("16", []),
            "K.ZOLI/12RAMPA": ("12", ["K. Zoli"]),
            "20.ramp": ("20", []),
            "15/rampa/k.zoli": ("15", ["K. Zoli"]),
            "Pót a 24 Rámpán": ("", []),
            "23 RAMPAN A POT": ("23", []),
            "15 ramp / U.Krisz": ("15", ["U. Krisz"]),
            "G-Lokációra kikészítve / 20 RAMP - TAMÁS": ("20", ["Tamás"]),
            "16os kikészítés": ("", []),
            "17-esen": ("", []),
            "PLOMBA008456": ("", []),
            "TEV6176195/MO": ("", []),
            "1 colli továbbengedett": ("", []),
            "Folyamatban": ("", []),
        }

        for note, expected in cases.items():
            with self.subTest(note=note):
                meta = app._tv_extract_loading_note_meta(note)
                self.assertEqual(meta["ramp"], expected[0])
                self.assertEqual(meta["loaders"], expected[1])

    def test_tv_ops_report_lists_only_fixed_expected_loadings(self):
        def awb(status, ready, note="", lmp="TEMU"):
            return {
                "ecomm_status": status,
                "status": status,
                "is_issued": False,
                "is_ready": ready,
                "lmp": lmp,
                "load_note": note,
            }

        state = {
            "status": "ready",
            "last_refresh": datetime(2026, 7, 2, 9, 30),
            "df": pd.DataFrame(),
            "outbound_cards": [
                {
                    "plate": "FIX-111",
                    "status_label": "Rakodas",
                    "checkin": datetime(2026, 7, 2, 8, 45),
                    "in_rest": False,
                    "glabs_items": [
                        {
                            "glabs_id": "ALL-READY",
                            "awb_items": [
                                awb("Kiadható", True, "16 RAMP / Gazsó"),
                                awb("Kiadható", True, "Gazsó"),
                            ],
                        },
                        {
                            "glabs_id": "HALF-NEAR",
                            "awb_items": [
                                awb("Kiadható", True, "20 RAMP / TAMÁS"),
                                awb("Kiadható", True, "20 RAMP / TAMÁS"),
                                awb("Megérkezett", False, "20 RAMP / TAMÁS"),
                                awb("Megérkezett", False, "20 RAMP / TAMÁS"),
                            ],
                        },
                        {
                            "glabs_id": "HALF-FAR",
                            "awb_items": [
                                awb("Kiadható", True),
                                awb("Kiadható", True),
                                awb("Felvéve", False),
                                awb("Értesítő", False),
                            ],
                        },
                    ],
                },
                {
                    "plate": "FIX-222",
                    "status_label": "Rakodas",
                    "checkin": None,
                    "in_rest": False,
                    "glabs_items": [{
                        "glabs_id": "READY-NO-DRIVER",
                        "awb_items": [
                            awb("Kiadható", True, "12 RAMP / K.ZOLI"),
                            awb("Kiadható", True, "K.ZOLI"),
                        ],
                    }],
                },
            ],
            "kpi": {
                "shift_name": "Nappali muszak",
                "shift_range": "08:00-20:00",
                "shift_elapsed_h": 2.0,
            },
        }

        with patch.object(app.data_cache, "get_state", return_value=state):
            payload = app.server.test_client().get("/api/tv/ops").get_json()

        fixed = payload["report"]["fixed_loadings"]
        by_glabs = {row["glabs_id"]: row for row in fixed}
        self.assertIn("ALL-READY", by_glabs)
        self.assertIn("HALF-NEAR", by_glabs)
        self.assertIn("READY-NO-DRIVER", by_glabs)
        self.assertNotIn("HALF-FAR", by_glabs)
        self.assertEqual(by_glabs["ALL-READY"]["ramp"], "16")
        self.assertEqual(by_glabs["ALL-READY"]["loader"], "Gazsó")
        self.assertEqual(by_glabs["HALF-NEAR"]["reason"], "fél felett + közel")
        self.assertEqual(by_glabs["READY-NO-DRIVER"]["loader"], "K. Zoli")

    def test_tv_ops_report_does_not_truncate_fixed_expected_loadings(self):
        cards = []
        for idx in range(10):
            cards.append({
                "plate": f"FIX-{idx:03d}",
                "status_label": "Rakodas",
                "checkin": datetime(2026, 7, 2, 8, 45),
                "in_rest": False,
                "glabs_items": [{
                    "glabs_id": f"READY-{idx:03d}",
                    "awb_items": [
                        {
                            "ecomm_status": "KiadhatĂł",
                            "status": "KiadhatĂł",
                            "is_issued": False,
                            "is_ready": True,
                            "lmp": "TEMU",
                            "load_note": "16 RAMP / GazsĂł",
                        },
                    ],
                }],
            })

        state = {
            "status": "ready",
            "last_refresh": datetime(2026, 7, 2, 9, 30),
            "df": pd.DataFrame(),
            "outbound_cards": cards,
            "kpi": {
                "shift_name": "Nappali muszak",
                "shift_range": "08:00-20:00",
                "shift_elapsed_h": 2.0,
            },
        }

        with patch.object(app.data_cache, "get_state", return_value=state):
            payload = app.server.test_client().get("/api/tv/ops").get_json()

        self.assertEqual(len(payload["report"]["fixed_loadings"]), 10)

    def test_tv_ops_truck_cards_do_not_expose_weight_colli_or_pallet_metrics(self):
        state = {
            "status": "ready",
            "last_refresh": datetime(2026, 7, 2, 9, 30),
            "df": pd.DataFrame(),
            "outbound_cards": [{
                "plate": "ABC-123",
                "status_label": "Rakodas",
                "checkin": datetime(2026, 7, 2, 8, 45),
                "in_rest": False,
                "glabs_items": [{
                    "glabs_id": "GL-1",
                    "awb_items": [{
                        "ecomm_status": "Felvéve",
                        "is_issued": False,
                        "is_ready": True,
                        "lmp": "LMP1",
                        "boxes": 11,
                        "weight": 2222,
                        "pallets_issued": 3,
                        "load_note": "K.ZOLI/12RAMPA",
                    }],
                }],
            }],
            "kpi": {
                "shift_kg": 1200,
                "shift_count": 2,
                "shift_name": "Nappali muszak",
                "shift_range": "08:00-20:00",
                "shift_elapsed_h": 2.0,
            },
        }

        with patch.object(app.data_cache, "get_state", return_value=state):
            payload = app.server.test_client().get("/api/tv/ops").get_json()

        truck = payload["trucks"][0]
        self.assertEqual(truck["plate"], "ABC-123")
        self.assertEqual(truck["ready"], 1)
        self.assertEqual(truck["total"], 1)
        self.assertNotIn("boxes", truck)
        self.assertNotIn("weight", truck)
        self.assertNotIn("pallets", truck)
        self.assertNotIn("colli", truck)
        self.assertEqual(truck["ramp"], "12")
        self.assertEqual(truck["loader"], "K. Zoli")
        self.assertEqual(truck["units"], [{"status": "Felvéve", "ready": True}])

    def test_tv_version_endpoint_returns_update_signature(self):
        response = app.server.test_client().get("/api/tv/version")
        payload = response.get_json()

        self.assertEqual(response.status_code, 200)
        self.assertTrue(payload["ok"])
        self.assertTrue(payload["signature"])
        self.assertTrue(payload["asset_signature"])
        self.assertTrue(payload["server_signature"])
        self.assertGreaterEqual(payload["asset_count"], 2)
        self.assertGreaterEqual(payload["server_count"], 1)
        self.assertIn("no-cache", response.headers.get("Cache-Control", ""))

    def test_header_saturation_uses_tracking_kpi_values(self):
        style, title, hover = app._header_saturation_style_and_title({
            "tracking": {
                "current_warehouse_saturation": {"kg": 139744, "ratio": 0.279488, "capacity": 500000},
            }
        })

        self.assertEqual(style["--warehouse-sat-pct"], "27.95%")
        self.assertIn("139 744 kg / 500 000 kg", title)
        self.assertIn("28%", hover)
        self.assertIn("139 744 kg", hover)

    def test_inbound_stat_counts_follow_card_filter_tokens(self):
        df = pd.DataFrame([
            {
                "awb": "A1",
                "is_stored": False,
                "is_at_hu": True,
                "priority_group": app.GRP_DRIVER_ACTIVE_HIGH,
                "is_shippable": True,
            },
            {
                "awb": "A2",
                "is_stored": False,
                "is_at_hu": False,
                "priority_group": app.GRP_DRIVER_PASSIVE_LOW,
                "is_shippable": False,
            },
            {
                "awb": "A3",
                "is_stored": True,
                "is_at_hu": True,
                "priority_group": app.GRP_DRIVER_ACTIVE_HIGH,
                "is_shippable": True,
            },
        ])

        original = app.storage_manager.get_snapshot_rows
        app.storage_manager.get_snapshot_rows = lambda _exclude: [
            {"awb": "OLD", "am_time": datetime.now()},
        ]
        try:
            counts = app._inbound_stat_counts({"df": df})
        finally:
            app.storage_manager.get_snapshot_rows = original

        self.assertEqual(counts["all"], "2")
        self.assertEqual(counts["at_hu"], "1")
        self.assertEqual(counts["driver"], "1")
        self.assertEqual(counts["scheduled"], "1")
        self.assertEqual(counts["shippable"], "1")
        self.assertEqual(counts["betarolt"], "2")

    def test_tracking_kpis_match_kpi_tracking_workbook_logic(self):
        now = datetime(2026, 6, 15, 10, 0)
        epoch = datetime(1899, 12, 30)

        def serial(dt):
            return (dt - epoch).total_seconds() / 86400

        raw = pd.DataFrame([
            {"awb": "172-03095993", "weight": 100, "boxes": 5, "parcel_count": 50, "status_n": "Felvéve", "ak_raw": serial(now), "expected_arrival_raw": serial(now)},
            {"awb": "172-03095993-1", "weight": 50, "boxes": 2, "parcel_count": 20, "status_n": "Kiadható", "ak_raw": serial(now), "expected_arrival_raw": serial(now)},
            {"awb": "555-11111111", "weight": 25, "boxes": 1, "parcel_count": 10, "status_n": "Mitörténik", "ak_raw": serial(now), "expected_arrival_raw": serial(now)},
            {"awb": "111-00000000", "weight": 30, "boxes": 1, "parcel_count": 15, "status_n": "Értesítő", "ak_raw": serial(now - timedelta(hours=13)), "expected_arrival_raw": serial(now - timedelta(days=1))},
            {"awb": "222-00000000", "weight": 40, "boxes": 1, "parcel_count": 10, "status_n": "Megérkezett", "ak_raw": serial(now - timedelta(hours=1)), "expected_arrival_raw": serial(now - timedelta(days=1))},
            {"awb": "333-00000000", "weight": 20, "boxes": 1, "parcel_count": 5, "status_n": "Szemlés", "ak_raw": serial(now - timedelta(hours=13)), "expected_arrival_raw": serial(now - timedelta(days=1))},
            {"awb": "444-00000000", "weight": 70, "boxes": 7, "parcel_count": 50, "status_n": "Elindult", "ak_raw": serial(now), "expected_arrival_raw": serial(now + timedelta(days=1))},
            {"awb": "666-00000000", "weight": 90, "boxes": 3, "parcel_count": 60, "status_n": "Kiadva", "ak_raw": serial(now), "expected_arrival_raw": serial(now)},
        ])

        tracking = data_reader._compute_tracking_kpis(raw, now, epoch)

        self.assertEqual(tracking["current_warehouse_saturation"]["kg"], 175)
        self.assertAlmostEqual(tracking["current_warehouse_saturation"]["ratio"], 0.00035)
        self.assertEqual(tracking["to_transfer_12h_kg"], 90)
        self.assertEqual(tracking["to_release_12h_kg"], 75)
        self.assertEqual(tracking["ata_more_than_12h_not_transferred_kg"], 50)
        self.assertAlmostEqual(tracking["kg_per_parcel"], 425 / 220)
        self.assertEqual(tracking["kg_per_colli"], 30)
        self.assertEqual(tracking["active_warehouse_items"], 2)

        by_date = {row["date"]: row for row in tracking["expected_arrivals"]}
        today = by_date["2026-06-15"]
        tomorrow = by_date["2026-06-16"]
        self.assertEqual(today["expected_arrivals_kg"], 265)
        self.assertEqual(today["already_arrived_kg"], 265)
        self.assertEqual(today["already_transferred_kg"], 265)
        self.assertEqual(tomorrow["remaining_expected_arrivals_kg"], 70)
        self.assertAlmostEqual(tomorrow["remaining_expected_parcel_kg"], 50 * (425 / 220))
        self.assertEqual(tomorrow["remaining_expected_colli_kg"], 210)

    def test_tracking_kpis_keep_workbook_helper_rows_in_unit_basis(self):
        now = datetime(2026, 6, 15, 10, 0)
        epoch = datetime(1899, 12, 30)
        raw = pd.DataFrame([
            {"awb": "172-03095993", "weight": 100, "boxes": 5, "parcel_count": 10, "status_n": "felveve"},
            {"awb": "", "weight": 40, "boxes": 0, "parcel_count": 0, "status_n": ""},
        ])

        tracking = data_reader._compute_tracking_kpis(raw, now, epoch)

        self.assertEqual(tracking["current_warehouse_saturation"]["kg"], 100)
        self.assertEqual(tracking["active_warehouse_items"], 1)
        self.assertEqual(tracking["kg_per_parcel"], 14)

    def test_ecomm_operational_read_range_covers_current_workbook_growth(self):
        self.assertEqual(data_reader._ECOMM_MAX_EXCEL_ROW, 25000)
        self.assertEqual(data_reader._ecomm_data_nrows(12), 24988)

    def test_truck_color_map_uses_far_apart_hues_for_common_count(self):
        colors = app._make_id_color_map([f"TRUCK-{i}" for i in range(6)])
        hues = []
        for color in colors.values():
            self.assertTrue(color.startswith("hsl("), color)
            hues.append(float(color[4:].split(",", 1)[0]))

        for i, hue in enumerate(hues):
            for other in hues[i + 1:]:
                self.assertGreaterEqual(app._hue_distance(hue, other), 40)

    def test_truck_color_map_does_not_repeat_red_family_in_first_palette_cycle(self):
        colors = app._make_id_color_map([f"TRUCK-{i}" for i in range(24)])
        hues = [float(color[4:].split(",", 1)[0]) for color in colors.values()]
        red_family = [h for h in hues if app._is_red_family_hue(h)]

        self.assertEqual(red_family, [8.0])

    def test_truck_color_map_keeps_similar_plate_numbers_apart(self):
        colors = app._make_id_color_map(["SVT-918", "SVT-912"])
        hue_918 = float(colors["SVT-918"][4:].split(",", 1)[0])
        hue_912 = float(colors["SVT-912"][4:].split(",", 1)[0])

        self.assertGreaterEqual(app._hue_distance(hue_918, hue_912), 40)

    def test_active_truck_palette_rows_match_sidebar_scope(self):
        df = pd.DataFrame([
            {"awb": "A1", "rendszam": "AAA-111", "is_stored": False, "is_en_route": False},
            {"awb": "A2", "rendszam": "BBB-222", "is_stored": True, "is_en_route": False},
            {"awb": "A3", "rendszam": "CCC-333", "is_stored": False, "is_en_route": True},
        ])

        rows = app._active_truck_palette_rows(df)
        colors = app._truck_plate_color_map(rows)

        self.assertEqual([row["rendszam"] for row in rows], ["AAA-111"])
        self.assertIn("AAA-111", colors)
        self.assertNotIn("BBB-222", colors)
        self.assertNotIn("CCC-333", colors)

    def test_truck_groups_include_plt_only_and_en_route_trucks(self):
        now = datetime.now()
        rows = [
            # PLT-only truck (no ULDs) — used to be filtered out entirely.
            {"awb": "P1", "rendszam": "SYE-225", "cargo_type": "PLT", "uld_number": "",
             "is_en_route": False, "is_shippable": True, "lmp": "TEMU", "am_time": now},
            # ULD truck.
            {"awb": "U1", "rendszam": "ULD-1", "cargo_type": "ULD", "uld_number": "AKE111",
             "is_en_route": False, "is_shippable": False, "lmp": "AT/HU", "am_time": now},
            # En-route truck (all items still arriving).
            {"awb": "R1", "rendszam": "ROUTE-9", "cargo_type": "PLT", "uld_number": "",
             "is_en_route": True, "is_shippable": False, "lmp": "4PX", "am_time": now,
             "truck_arrival_time": now + timedelta(minutes=20)},
        ]
        colors = app._truck_plate_color_map(rows)
        groups = {g["plate"]: g for g in app._truck_groups(rows, colors)}

        self.assertIn("SYE-225", groups)            # PLT-only now listed
        self.assertIn("ULD-1", groups)
        self.assertIn("ROUTE-9", groups)
        self.assertTrue(groups["SYE-225"]["has_plt"])
        self.assertEqual(len(groups["ULD-1"]["ulds"]), 1)
        self.assertFalse(groups["SYE-225"]["is_en_route_truck"])
        self.assertTrue(groups["ROUTE-9"]["is_en_route_truck"])


if __name__ == "__main__":
    unittest.main()
