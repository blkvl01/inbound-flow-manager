# -*- coding: utf-8 -*-
"""Generate Flow Manager INBOUND user guide."""
import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

HERE = Path(__file__).parent
OUTPUT = HERE / "Flow_Manager_Inbound_leiras.docx"


def setup_styles(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    ts = doc.styles["Title"]
    ts.font.name = "Calibri"
    ts.font.size = Pt(26)
    ts.font.bold = True
    ts.font.color.rgb = RGBColor(0x1A, 0x5C, 0x6E)
    ts.paragraph_format.space_after = Pt(4)

    ss = doc.styles["Subtitle"]
    ss.font.name = "Calibri"
    ss.font.size = Pt(13)
    ss.font.italic = True
    ss.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    ss.paragraph_format.space_after = Pt(20)

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0x1A, 0x5C, 0x6E)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(8)

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0x2D, 0x8C, 0x6E)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)


def add_img(doc, name, caption, width=None):
    p = HERE / name
    if not p.exists():
        doc.add_paragraph(f"[Kép: {name}]").alignment = WD_ALIGN_PARAGRAPH.CENTER
        return
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(p), width=width or Inches(6.2))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    cap.paragraph_format.space_after = Pt(12)


def add_tip(doc, text):
    p = doc.add_paragraph()
    r = p.add_run("ℹ️  " + text)
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x44, 0x7A, 0x44)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(10)


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix + " ")
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)


def build():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    setup_styles(doc)

    # Title page
    for _ in range(4):
        doc.add_paragraph()

    title = doc.add_paragraph("Flow Manager", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph(
        "INBOUND nézet – Felhasználói útmutató",
        style="Subtitle",
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = info.add_run("HGL Group Hungary – Ecommerce Operations")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.add_page_break()

    # --- BEVEZETES ---
    doc.add_heading("Mire való a Flow Manager?", level=1)

    doc.add_paragraph(
        "A Flow Manager egy belső dashboard, ami a raktári csapat munkáját segíti "
        "azzal, hogy valós időben mutatja az aktív bejövő szállítmányokat. "
        "Nem kell keresgélni az Excelben, nem kell telefonálgatni – egyetlen "
        "képernyőn látod, mi van a raktárban, mi jön, és melyik szállítmányt "
        "érdemes először kezelni."
    )

    doc.add_paragraph(
        "Az alkalmazás automatikusan olvassa az E_COMM és BUD-Pallets Excel fájlokat, "
        "és ezekből építi fel a kártyákat, amiket a következő oldalakon részletezünk. "
        "A háttérben 10 percenként frissül az adat, de ha az Excel fájl változik, "
        "azt azonnal észreveszi."
    )

    add_tip(doc, "A Flow Manager nem módosítja az Excel fájlokat – csak olvassa őket.")

    # --- DASHBOARD ---
    doc.add_heading("A Dashboard felépítése", level=1)

    doc.add_paragraph(
        "Amikor megnyitod a Flow Managert, az alábbi képernyő fogad. "
        "Nézzük végig, mit látsz rajta felülről lefele:"
    )

    add_img(
        doc,
        "img_dashboard.png",
        "A Flow Manager főképernyője – INBOUND nézet, aktív szállítmányok kártyái",
    )

    # --- FEJLEC ---
    doc.add_heading("Fejléc és KPI sáv", level=2)

    add_img(
        doc,
        "img_header.png",
        "A fejléc sáv: műszak kg, várható beérkező, és frissítés ideje",
        width=Inches(5.8),
    )

    doc.add_paragraph("A felső sávban három fontos információ található:")

    add_bullet(
        doc,
        "– az aktuális műszakban felvett összsúly kilogrammban.",
        bold_prefix="Nappali/Éjszakai Műszak",
    )
    add_bullet(
        doc,
        "– a következő órákban várhatóan beérkező szállítmányok "
        "összsúlya. Mellette színes sáv mutatja az összetelt: "
        "Felvéve (zöld), Értesítő (narancssárga), Megérkezett (kék) bontásban.",
        bold_prefix="Várható Beérkező",
    )
    add_bullet(
        doc,
        "– a jobb sarokban látod, mikor frissült utoljára az adat, "
        "és mikor lesz a következő frissítés.",
        bold_prefix="Frissítés ideje",
    )

    add_tip(doc, "A Várható Beérkező KPI-ra kattintva részletes bontás nyílik meg.")

    # --- SZURO GOMBOK ---
    doc.add_heading("Szűrő gombok és jelmagyarázat", level=2)

    add_img(
        doc,
        "img_stats.png",
        "A statisztikai szűrő gombok – kattintással szűrheted a kártyákat",
        width=Inches(5.2),
    )

    doc.add_paragraph(
        "A fejléc alatt egy sor szűrő gomb található. Mindegyik gomb tetején "
        "egy szám mutatja, hány tétel tartozik az adott kategóriába. "
        "Ha rákattintasz egy gombra, csak az ahhoz tartozó kártyákat látod:"
    )

    add_bullet(doc, "– az összes aktív szállítmány száma.", bold_prefix="Aktív tétel")
    add_bullet(
        doc,
        "– osztrák és magyar viszonylatok, amik prioritást élveznek.",
        bold_prefix="AT/HU",
    )
    add_bullet(
        doc,
        "– azoknál a tételeknél, ahol a sofőr már bejelentkezett a telephelyen.",
        bold_prefix="Sofőr helyszínen",
    )
    add_bullet(
        doc,
        "– a sofőr még nem érkezett meg, vagy éppen pihenőn van.",
        bold_prefix="Nincs itt / pihenőn",
    )
    add_bullet(
        doc,
        "– a BUD-Pallets szerint kiadásra kész tételek.",
        bold_prefix="Kiadható",
    )
    add_bullet(
        doc,
        "– azok a tételek, amiket valaki manuálisan betároltnak jelölt.",
        bold_prefix="Betárolt tételek",
    )

    doc.add_paragraph(
        "A gombok alatt egy színes jelmagyarázat segít eligazodni a kártyák közti "
        "sorrendben. A különböző színek és kategóriák jelzik, melyik szállítmány "
        "miért van előrébb vagy hátrébb a listában."
    )

    # --- KARTYAK ---
    doc.add_heading("A kártyák – mit mutatnak?", level=1)

    doc.add_paragraph(
        "Minden aktív szállítmány egy-egy kártyaként jelenik meg. "
        "A kártyák nem véletlenszerűen vannak – a rendszer automatikusan rangsorolja "
        "őket fontosság szerint, úgyhogy ami elöl van, azt érdemes először kezelni."
    )

    doc.add_heading("Egy kártya részletesen", level=2)

    add_img(
        doc,
        "img_card_full.png",
        "Egy teljes kártya: kategória, AWB, rakodás tételei, és sofőr státusz",
        width=Inches(2.8),
    )

    doc.add_paragraph(
        "Fentről lefele haladva ezeket az információkat találod egy kártyán:"
    )

    doc.add_heading("Kategória fejléc", level=2)

    doc.add_paragraph(
        "A kártya tetején színes sáv mutatja, milyen kategóriába esik a szállítmány. "
        "Például: \"Sofőr helyszínen – sok kiadható\" azt jelenti, hogy a sofőr "
        "már a telephelyen van, és az adott rakodásban több tétel is kiadásra kész. "
        "Ez a legmagasabb prioritás – ezeket kell először kezelni."
    )

    doc.add_paragraph(
        "A kategória mellett gombok is lehetnek: a \"Betárolva\" gomb (zöld) a manuális "
        "betárolásra szolgál, a \"Kiadható\" címke (türkiz) jelzi, ha a tétel kiadásra kész, "
        "és az \"ULD\" jelölés mutatja, ha a szállítmány konténerben érkezett (nem palettán)."
    )

    doc.add_heading("AWB és szállítmány adatok", level=2)

    doc.add_paragraph(
        "A nagy, fehér szám az AWB (légifuvarlevl) szám – ez azonosítja a szállítmányt. "
        "Mellette halványabb betűkkel az ULD/konténer szám látszik, ha van."
    )

    doc.add_paragraph(
        "Alatta a fuvarozó neve (pl. TEMU RO CARGUS), "
        "a colli szám (csomagok száma) és az összsúly kilogrammban. "
        "A \"Felvéve\" időpont mutatja, mikor lett a szállítmány felvéve a rendszerbe, "
        "és alatta zölddel jelenik meg, mennyi ideje várakozik."
    )

    doc.add_heading("Rakodás állapota (GLABS)", level=2)

    doc.add_paragraph(
        "Ha a szállítmány egy aktív rakodáshoz tartozik, megjelenik egy progress sáv "
        "a GLABS azonosító mellett. Ez mutatja, hány tétel kiadható a rakodásból "
        "az összeshez képest (pl. \"5/5 kiadható\")."
    )

    doc.add_paragraph(
        "A \"Rakodás tételei\" lenyíló listában látod az adott rakodás összes tételét – "
        "nem csak az aktuális kártyáét, hanem a rakodás többi szállítmányát is. "
        "Így egy pillantással láthatod, hogy a kamion többi tétele hogyan áll."
    )

    doc.add_heading("Sofőr státusz", level=2)

    doc.add_paragraph(
        "A kártya alján egy külön blokk mutatja a sofőr jelenlétét. "
        "Ha a sofőr bejelentkezett, zöld ponttal és a bejelentkezs "
        "időpontjával jelenik meg. Alatta látod, mennyi ideje várakozik."
    )

    doc.add_paragraph(
        "Ha a sofőr pihenőn van, azt sárga/narancssárga jelzés mutatja. "
        "Ha nincs bejelentkezve, szürke háttérrel jelzi a rendszer."
    )

    # --- SORREND ---
    doc.add_heading("A kártyák sorrendje – mi alapján rangsorol?", level=1)

    doc.add_paragraph(
        "A Flow Manager nem csak listázza a szállítmányokat, hanem automatikusan "
        "sorba rakja őket aszerint, melyiket érdemes először kezelni. "
        "A rangsorolás több tényezőt vesz figyelembe:"
    )

    add_img(
        doc,
        "img_cards_top_row.png",
        "Az első sor kártyák: \"Sofőr helyszínen – sok kiadható\" kategória",
        width=Inches(6.0),
    )

    doc.add_paragraph(
        "A legfontosabb szempont, hogy a sofőr ott van-e a telephelyen és hány "
        "kiadásra kész tétel van a rakodásban. Az alábbi sorrend érvényesül:"
    )

    add_bullet(
        doc,
        "A sofőr a telephelyen van, és a rakodásban sok tétel kiadható. "
        "Ezek a kártyák narancssárga/arany szegéllyel jelennek meg.",
        bold_prefix="1. Sofőr itt + sok kiadható –",
    )
    add_bullet(
        doc,
        "A sofőr itt van, de kevesebb tétel kiadható.",
        bold_prefix="2. Sofőr itt + kevés kiadható –",
    )
    add_bullet(
        doc,
        "A szállítmány szerepel egy rakodásban, de a sofőr "
        "még nem érkezett meg vagy pihenőn van.",
        bold_prefix="3. Nincs itt / pihenőn –",
    )

    add_img(
        doc,
        "img_cards_bottom_row.png",
        "A második sor: \"Sofőr helyszínen – kevés kiadható\" kategória",
        width=Inches(6.0),
    )

    doc.add_paragraph(
        "Azonos kategórián belül az AT/HU viszonylatok előrébb kerülnek, "
        "majd a TEMU szállítmányok, és végül a többi fuvarozó. "
        "A várakozási idő is számít: ha két szállítmány ugyanabban a "
        "kategóriában van, a régebben várakozó kerül előrébb."
    )

    # --- BETAROLVA ---
    doc.add_heading("Betárolva gomb – manuális jelölés", level=1)

    doc.add_paragraph(
        "Minden kártyán van egy \"Betárolva\" gomb. Ha egy szállítmányt "
        "fizikailag betároltál a raktárba, kattints rá. Ilyenkor a kártya "
        "eltűnik az aktív nézetből (hiszen már nem kell foglalkozni vele), "
        "és a \"Betárolt tételek\" szűrővel tudod megnézni a már kezelt tételeket."
    )

    doc.add_paragraph(
        "A betárolás állapotát mindenki látja, aki a Flow Managert használja – "
        "közös fájlban tárolódik, szóval ha te betárolsz valamit, a kollégáid "
        "képernyőjén is eltűnik az a kártya."
    )

    add_tip(
        doc,
        "Ha tévedésből jelölted betároltnak, a Betárolt tételek "
        "szűrőben visszavonhatod a műveletet.",
    )

    doc.add_paragraph(
        "Ha egy tételt betárolsz, de a BUD-Pallets még nem mutatja kiadhatónak, "
        "a rendszer automatikusan hozzáadja a rakodás progresséhez – így a "
        "kapcsolódó kártyákon a kiadható arány is frissül."
    )

    # --- INBOUND OUTBOUND ---
    doc.add_heading("INBOUND és OUTBOUND nézet", level=1)

    doc.add_paragraph(
        "A jobb felső sarokban két fül látható: INBOUND és OUTBOUND. "
        "Ez a leírás az INBOUND nézetről szól – a bejövő szállítmányokról. "
        "Az OUTBOUND nézet a kimenő rakodásokat mutatja rendszám és kamion szerint, "
        "ami egy másik munkafolyamathoz tartozik."
    )

    add_tip(doc, "A nézet váltásakor a szűrők automatikusan visszaállnak.")

    # --- VILAGOS MOD ---
    doc.add_heading("Világos és sötét mód", level=1)

    doc.add_paragraph(
        "A fejléc jobb oldalán a \"Világos mód\" / \"Sötét mód\" gombbal válthatsz "
        "a két megjelenés között. Az alapértelmezett a sötét mód, "
        "de ha jobban szereted a világos hátteret, bátran válts."
    )

    # --- TIPPEK ---
    doc.add_heading("Gyakorlati tippek", level=1)

    add_bullet(
        doc,
        "Mindig a lista tetején lévő kártyákkal foglalkozz először – azok a legsürgősebbek.",
    )
    add_bullet(
        doc,
        "Ha a sofőr már a telephelyen van és sok tétel kiadható, az a legmagasabb prioritás.",
    )
    add_bullet(
        doc,
        "Ha betároltál egy tételt, jelöld be – így a többiek is látják, "
        "hogy már nem kell vele foglalkozni.",
    )
    add_bullet(
        doc,
        "A \"Rakodás tételei\" lenyílót nyisd ki, ha látni akarod, "
        "hogy a kamion többi szállítmánya hogyan áll.",
    )
    add_bullet(
        doc,
        "Ne aggódj, ha az oldal néha pár másodpercig \"Betöltés...\" "
        "feliratot mutat – az Excel fájl nagy, időbe telik beolvasni.",
    )
    add_bullet(
        doc,
        "Ha 30 percig nem használod a felületet, az alkalmazás automatikusan leáll. "
        "Ilyenkor csak indítsd újra.",
    )

    # Footer
    doc.add_paragraph()
    fl = doc.add_paragraph()
    fl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fl.add_run("— Flow Manager  ·  HGL Group Hungary  ·  Ecommerce Operations —")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

    doc.save(str(OUTPUT))
    print(f"Saved: {OUTPUT}")
    print(f"Size: {OUTPUT.stat().st_size:,} bytes")


if __name__ == "__main__":
    build()
